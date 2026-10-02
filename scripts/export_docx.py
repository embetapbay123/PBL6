"""Export the current PBL6 Markdown report set to readable Word copies."""
from __future__ import annotations

import os
import re
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "word"
ASSETS = Path(os.environ.get("PBL6_WORD_ASSETS", str(Path(os.environ["TEMP"]) / "pbl6-word-assets-2.1")))

SOURCE_NAMES = """
scope decisions srs functional-requirements business-rules state-transitions rbac
user-stories use-cases ui-flows architecture data-dictionary api-spec
integration-contract security-privacy ai-design ai-evaluation test-plan
requirements-acceptance traceability deployment project-plan
""".split()

TITLES = {
    "scope": "Phạm vi và thuật ngữ PBL6",
    "decisions": "Nhật ký quyết định PBL6",
    "srs": "Đặc tả yêu cầu phần mềm PBL6",
    "functional-requirements": "Danh mục yêu cầu chức năng và phi chức năng",
    "business-rules": "Quy tắc nghiệp vụ PBL6",
    "state-transitions": "Vòng đời và chuyển trạng thái nghiệp vụ",
    "rbac": "Ma trận quyền truy cập PBL6",
    "user-stories": "Câu chuyện người dùng PBL6",
    "use-cases": "Đặc tả Use Case PBL6",
    "ui-flows": "Luồng giao diện và danh mục màn hình",
    "architecture": "Kiến trúc hệ thống PBL6",
    "data-dictionary": "Từ điển dữ liệu PBL6",
    "api-spec": "Đặc tả API PBL6",
    "integration-contract": "Hợp đồng tích hợp giữa các service",
    "security-privacy": "Bảo mật dữ liệu cá nhân và audit",
    "ai-design": "Thiết kế chatbot và recommendation",
    "ai-evaluation": "Kế hoạch đánh giá AI",
    "test-plan": "Kế hoạch và trường hợp kiểm thử",
    "requirements-acceptance": "Tiêu chí nghiệm thu theo yêu cầu",
    "traceability": "Ma trận truy vết yêu cầu",
    "deployment": "Hướng dẫn triển khai và demo",
    "project-plan": "Kế hoạch bàn giao tài liệu PBL6",
    "diagrams": "Sơ đồ và hướng dẫn đọc PBL6",
}

LANDSCAPE = {
    "functional-requirements", "rbac", "user-stories", "use-cases", "data-dictionary",
    "requirements-acceptance", "traceability", "test-plan", "diagrams",
}
RECORD_TABLES = {"functional-requirements", "user-stories", "requirements-acceptance", "test-plan"}


def font_run(run, name="Arial", size=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def cell_padding(cell, top=85, left=100, bottom=85, right=100):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for name, amount in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        item = margins.find(qn(f"w:{name}"))
        if item is None:
            item = OxmlElement(f"w:{name}")
            margins.append(item)
        item.set(qn("w:w"), str(amount))
        item.set(qn("w:type"), "dxa")


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    trpr.append(flag)


def keep_row_together(row):
    trpr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    trpr.append(flag)


def setup_doc(stem):
    doc = Document()
    section = doc.sections[0]
    if stem == "diagrams":
        section.page_width, section.page_height = Cm(42), Cm(29.7)
        section.left_margin = section.right_margin = Cm(1.8)
    elif stem in LANDSCAPE:
        section.page_width, section.page_height = Cm(29.7), Cm(21)
        section.left_margin = section.right_margin = Cm(1.65)
    else:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(1.75)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)

    styles = doc.styles
    body = styles["Normal"]
    body.font.name = "Arial"
    body.font.size = Pt(10 if stem in LANDSCAPE else 10.5)
    body.font.color.rgb = RGBColor(0, 0, 0)
    body.paragraph_format.space_after = Pt(5)
    body.paragraph_format.line_spacing = 1.13
    for name, size, before, after in (("Title", 20, 0, 16), ("Heading 1", 15, 16, 7),
                                       ("Heading 2", 12.5, 12, 5), ("Heading 3", 10.5, 9, 3)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # The default Word template gives Title a blue bottom border. Keep report
    # headings neutral and remove that inherited rule explicitly.
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    p = doc.add_paragraph(style="Title")
    p.add_run(TITLES[stem])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Phiên bản 2.1 Draft")
    font_run(r, size=10, color=(75, 85, 96))
    doc.core_properties.title = TITLES[stem]
    doc.core_properties.subject = "PBL6 Marketplace documentation 2.1 Draft"
    return doc


def make_target(src_file: Path, target: str) -> str:
    target = unquote(target)
    if target.startswith(("https://", "http://", "mailto:")):
        return target
    target_path, _, fragment = target.partition("#")
    if not target_path:
        return ""
    resolved = (src_file.parent / target_path).resolve()
    if resolved.parent == DOCS.resolve() and resolved.suffix.lower() == ".md" and resolved.stem in SOURCE_NAMES:
        return resolved.stem + ".docx"
    if "diagrams" in resolved.parts and resolved.suffix.lower() in (".md", ".puml", ".mmd", ".svg"):
        return "diagrams.docx"
    return str(resolved)


def hyperlink(paragraph, label: str, target: str):
    if not target:
        paragraph.add_run(label)
        return
    rel = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    pr.append(color)
    run.append(pr)
    t = OxmlElement("w:t")
    t.text = label
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, source: Path):
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()].replace("\\|", "|"))
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            font_run(run, "Consolas", 9)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            hyperlink(paragraph, link.group(1), make_target(source, link.group(2)))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:].replace("\\|", "|"))


def clean(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "").replace("\\|", "|").strip()


def table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def add_labeled(doc, label: str, value: str, source: Path, *, small=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2 if small else 4)
    run = p.add_run(label + ": ")
    run.bold = True
    add_inline(p, value, source)
    return p


def add_record_table(doc, headers, rows, source, stem):
    for cells in rows:
        if not cells or not cells[0]:
            continue
        title = clean(cells[0])
        p = doc.add_paragraph(style="Heading 3")
        p.paragraph_format.space_before = Pt(9)
        p.add_run(title)
        if stem == "user-stories" and len(cells) == 8:
            add_labeled(doc, "Vai trò", f"{cells[2]}   •   Ưu tiên: {cells[5]}   •   FR: {cells[6]}", source, small=True)
            add_labeled(doc, "Mong muốn", cells[3], source, small=True)
            add_labeled(doc, "Lợi ích", cells[4], source, small=True)
            add_labeled(doc, "Acceptance criteria", cells[7], source)
        elif stem == "requirements-acceptance" and len(cells) == 6:
            add_labeled(doc, "Yêu cầu", cells[1], source, small=True)
            for label, value in zip(("Given", "When", "Then", "Trạng thái"), cells[2:]):
                add_labeled(doc, label, value, source, small=label != "Then")
        elif stem == "functional-requirements":
            for label, value in zip(headers[1:], cells[1:]):
                add_labeled(doc, clean(label), value, source, small=len(value) < 80)
        elif stem == "test-plan":
            for label, value in zip(headers[1:], cells[1:]):
                add_labeled(doc, clean(label), value, source, small=len(value) < 80)
        else:
            for label, value in zip(headers[1:], cells[1:]):
                add_labeled(doc, clean(label), value, source, small=len(value) < 80)


def add_grid_table(doc, headers, rows, source, stem):
    n = len(headers)
    if n == 0:
        return
    table = doc.add_table(rows=1, cols=n)
    table.autofit = False
    section = doc.sections[0]
    width_inches = (section.page_width - section.left_margin - section.right_margin) / Inches(1)
    samples = [headers] + rows[:30]
    weights = []
    for j in range(n):
        lengths = [len(clean(row[j])) for row in samples if j < len(row)]
        typical = sorted(lengths)[min(len(lengths)-1, max(0, int(len(lengths)*0.65)))] if lengths else 10
        weights.append(max(8, min(44, typical)))
    total = sum(weights)
    widths = [width_inches * w / total for w in weights]
    for j, head in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.width = Inches(widths[j])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "24466B")
        cell_borders(cell)
        cell_padding(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(clean(head))
        font_run(r, size=8.4 if n >= 6 else 9, color=(255, 255, 255))
        r.bold = True
    repeat_header(table.rows[0])
    keep_row_together(table.rows[0])
    for i, values in enumerate(rows):
        if len(values) != n:
            continue
        row = table.add_row()
        keep_row_together(row)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_borders(cell)
            cell_padding(cell)
            if i % 2:
                shade(cell, "F3F7FB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if j == 0 and len(value) < 45:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value, source)
            for run in p.runs:
                font_run(run, size=8.3 if n >= 6 else 8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, lines, source, stem):
    parsed = [table_cells(line) for line in lines]
    if len(parsed) < 2:
        return
    headers = parsed[0]
    rows = [row for row in parsed[1:] if not all(re.fullmatch(r":?-{2,}:?", cell or "") for cell in row)]
    if stem in RECORD_TABLES and len(headers) >= 5:
        add_record_table(doc, headers, rows, source, stem)
    else:
        add_grid_table(doc, headers, rows, source, stem)


def add_image(doc, markdown_line, source: Path, stem: str):
    match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", markdown_line.strip())
    if not match:
        return
    label, target = match.groups()
    source_image = (source.parent / unquote(target)).resolve()
    if source_image.suffix.lower() == ".svg":
        relative = source_image.relative_to((DOCS / "diagrams").resolve())
        image = ASSETS / relative.with_suffix(".png")
    else:
        image = source_image
    if not image.exists():
        raise FileNotFoundError(image)
    section = doc.sections[0]
    max_width = (section.page_width - section.left_margin - section.right_margin) / Inches(1)
    # The illustrated handbook has a group heading, source heading, links and
    # caption on the same page. Leave enough room so Word does not orphan the
    # headings on an otherwise empty page.
    reserved_height = 2.1 if stem == "diagrams" else 0.65
    max_height = (section.page_height - section.top_margin - section.bottom_margin) / Inches(1) - reserved_height
    with Image.open(image) as pic:
        ratio = pic.width / pic.height
    width = min(max_width, max_height * ratio)
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image), width=Inches(width))
    if label:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(12)
        run = caption.add_run(label)
        font_run(run, size=9, color=(75, 85, 96))


def add_markdown(doc, source: Path, stem: str, *, combined=False):
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or re.match(r"^<a\s+id=", line):
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            if language == "mermaid":
                p = doc.add_paragraph()
                p.add_run("Sơ đồ context và container (mã Mermaid chỉnh sửa được trong bản Markdown)").italic = True
            for code_line in block:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.4)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                run = p.add_run(code_line or " ")
                font_run(run, "Consolas", 8.2)
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level == 1 and first_h1 and not combined:
                first_h1 = False
                i += 1
                continue
            first_h1 = False
            mapped = min(3, max(1, level if combined else level - 1))
            p = doc.add_paragraph(style=f"Heading {mapped}")
            add_inline(p, heading.group(2), source)
            i += 1
            continue
        if line.startswith("!["):
            add_image(doc, line, source, stem)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block, source, stem)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line), source)
            i += 1
            continue
        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.48)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.add_run(numbered.group(1) + ". ").bold = True
            add_inline(p, numbered.group(2), source)
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.3)
            add_inline(p, line.lstrip("> "), source)
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        if re.fullmatch(r"-{3,}", line.strip()):
            i += 1
            continue
        paragraph = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6})\s+|^\||^!\[|^```|^\s*[-*]\s+|^\s*\d+\.\s+|^>", lines[i]):
            paragraph.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph), source)


def export_one(stem):
    source = DOCS / (stem + ".md")
    doc = setup_doc(stem)
    add_markdown(doc, source, stem)
    output = OUT / (stem + ".docx")
    doc.save(output)
    return output


def export_diagrams():
    doc = setup_doc("diagrams")
    p = doc.add_paragraph()
    p.add_run("Tập hình này gom các chú giải ERD, sequence, activity và state. Mã nguồn .mmd/.puml vẫn nằm trong docs/diagrams để chỉnh sửa.")
    groups = [
        ("ERD", sorted((DOCS / "diagrams" / "erd").glob("[0-9][0-9]-*.md"))),
        ("Sequence", sorted((DOCS / "diagrams" / "sequence").glob("[0-9][0-9]-*.md"))),
        ("Activity", sorted((DOCS / "diagrams" / "activity").glob("[0-9][0-9]-*.md"))),
    ]
    for group, files in groups:
        doc.add_page_break()
        doc.add_paragraph(group, style="Heading 1")
        for position, source in enumerate(files):
            if position:
                doc.add_page_break()
            add_markdown(doc, source, "diagrams", combined=True)
    doc.add_page_break()
    doc.add_paragraph("State diagram", style="Heading 1")
    add_markdown(doc, DOCS / "diagrams" / "state" / "README.md", "diagrams", combined=True)
    output = OUT / "diagrams.docx"
    doc.save(output)
    return output


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    outputs = [export_one(stem) for stem in SOURCE_NAMES]
    outputs.append(export_diagrams())
    assert len(outputs) == 23
    for output in outputs:
        print(output.name, output.stat().st_size)


if __name__ == "__main__":
    main()
