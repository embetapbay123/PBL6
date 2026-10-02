"""Convert the five original PBL6 documents to auditable UTF-8 Markdown.

Requires python-docx and openpyxl. Source files are never modified.
Run from any directory. Imported source documents live in docs/legacy.
Existing imports are protected unless --force is explicitly supplied.
"""
from pathlib import Path
from datetime import date, datetime
from hashlib import sha256
from html import escape
from urllib.parse import quote
import json
import re
from zipfile import ZipFile
import sys

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs_old"
DEST = ROOT / "docs" / "legacy"
FILES = {
    "SRS_v1.2.docx": ("srs.md", "Đặc tả yêu cầu phần mềm"),
    "Use_Case_Specification (1).docx": ("use-cases.md", "Đặc tả ca sử dụng"),
    "Functional_Requirements.xlsm": ("functional-requirements.md", "Yêu cầu chức năng và thuộc tính chất lượng"),
    "User_Story.xlsm": ("user-stories.md", "User Story"),
    "Phân chia việc.xlsx": ("phan-cong.md", "Phân công công việc"),
}


def md(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    text = escape(str(value), quote=False).replace("\\", "\\\\")
    for char in ("|", "*", "_", "`", "[", "]"):
        text = text.replace(char, "\\" + char)
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>").replace("\t", "    ")


def table(rows):
    width = max(map(len, rows), default=0)
    if not width:
        return ""
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(r) + " |" for r in rows[1:])
    return "\n".join(lines)


def convert_word(path):
    document = Document(path)
    chunks = []
    paragraphs = tables = 0
    numbering = document.part.numbering_part.element
    abstract = {n.get(qn("w:abstractNumId")): n for n in numbering.findall(qn("w:abstractNum"))}
    numbers = {n.get(qn("w:numId")): n.find(qn("w:abstractNumId")).get(qn("w:val")) for n in numbering.findall(qn("w:num"))}
    counters = {}
    for block in document.element.body:
        if block.tag == qn("w:p"):
            p = Paragraph(block, document)
            if not p.text.strip():
                continue
            paragraphs += 1
            heading = re.match(r"Heading (\d+)", p.style.name)
            text = md(p.text)
            if p._p.pPr is None or p._p.pPr.numPr is None:
                counters.clear()
            if heading:
                text = "#" * min(int(heading.group(1)) + 1, 6) + " " + text
            elif p._p.pPr is not None and p._p.pPr.numPr is not None:
                props = p._p.pPr.numPr
                number_id = str(props.numId.val)
                level = str(props.ilvl.val if props.ilvl is not None else 0)
                definition = abstract[numbers[number_id]]
                lvl = next(n for n in definition.findall(qn("w:lvl")) if n.get(qn("w:ilvl")) == level)
                kind = lvl.find(qn("w:numFmt")).get(qn("w:val"))
                if kind == "decimal":
                    key = (number_id, level)
                    counters[key] = counters.get(key, int(lvl.find(qn("w:start")).get(qn("w:val"))) - 1) + 1
                    text = f"{counters[key]}. " + text
                else:
                    text = "- " + text
            chunks.append(text)
        elif block.tag == qn("w:tbl"):
            counters.clear()
            tables += 1
            t = Table(block, document)
            rows = []
            seen = set()
            for row in t.rows:
                values = []
                for cell in row.cells:
                    if cell._tc in seen:
                        values.append("")
                    else:
                        seen.add(cell._tc)
                        values.append(md(cell.text))
                rows.append(values)
            chunks.extend([f'<a id="table-{tables}"></a>', table(rows)])
    # These source documents contain no images; fail instead of silently dropping new media.
    with ZipFile(path) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        if media:
            raise ValueError(f"Images require explicit extraction: {path.name}: {media}")
    return "\n\n".join(chunks), {"nonempty_paragraphs": paragraphs, "tables": tables}


def convert_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=False)
    chunks = []
    stats = []
    for index, sheet in enumerate(workbook.worksheets, 1):
        chunks.extend([f'<a id="sheet-{index}"></a>', f"## {md(sheet.title)}"])
        chunks.append("Cột **Dòng Excel** giữ vị trí trong file gốc; các cột chữ cái tương ứng cột Excel. Ô trống được giữ nguyên, không tự điền dữ liệu từ dòng trên.")
        if sheet.merged_cells:
            chunks.append("Ô gộp trong nguồn: " + ", ".join(f"`{r}`" for r in sheet.merged_cells.ranges) + ". Nội dung nằm ở ô đầu của vùng gộp.")
        rows = [["Dòng Excel"] + [openpyxl.utils.get_column_letter(c) for c in range(1, sheet.max_column + 1)]]
        nonempty = formulas = comments = 0
        notes = []
        for row in sheet.iter_rows():
            if not any(c.value is not None for c in row):
                continue
            values = [f'<a id="sheet-{index}-row-{row[0].row}"></a>{row[0].row}']
            for cell in row:
                if cell.value is not None:
                    nonempty += 1
                formulas += cell.data_type == "f"
                value = md(cell.value)
                if cell.hyperlink:
                    target = cell.hyperlink.target or cell.hyperlink.location
                    value += f" (Liên kết: {md(target)})"
                if cell.comment:
                    comments += 1
                    notes.append(f"- {cell.coordinate}: {md(cell.comment.text)}")
                values.append(value)
            rows.append(values)
        chunks.append(table(rows))
        if notes:
            chunks.extend(["### Ghi chú ô", "\n".join(notes)])
        stats.append({"sheet": sheet.title, "state": sheet.sheet_state, "nonempty_rows": len(rows) - 1, "nonempty_cells": nonempty, "formulas": formulas, "comments": comments, "merged_ranges": [str(r) for r in sheet.merged_cells.ranges]})
    workbook.close()
    return "\n\n".join(chunks), {"sheets": stats}


def main():
    DEST.mkdir(exist_ok=True)
    force = "--force" in sys.argv[1:]
    existing = [output for output, _ in FILES.values() if (DEST / output).exists()]
    if existing and not force:
        raise SystemExit("Imports already exist in docs/legacy; use --force to replace them: " + ", ".join(existing))
    audit = []
    for filename, (output, title) in FILES.items():
        path = SOURCE / filename
        before = sha256(path.read_bytes()).hexdigest()
        body, stats = convert_word(path) if path.suffix == ".docx" else convert_workbook(path)
        header = f"# {title}\n\nNguồn: [{filename}](../../docs_old/{quote(filename)}).\n\n> Bản chuyển đổi nội dung từ tài liệu gốc, chưa hợp nhất các mâu thuẫn nghiệp vụ. Xem [báo cáo rà soát](../RA_SOAT.md) và [tài liệu 2.0 Draft](../README.md).\n\n"
        (DEST / output).write_text(header + body + "\n", encoding="utf-8")
        assert sha256(path.read_bytes()).hexdigest() == before
        audit.append({"source": filename, "output": output, "sha256": before, **stats})
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
